"""Smoke test of the office/PDF toolbox preinstalled for pi (QA round 1, PA-02).

Run inside the image as `python - < doc-tools-check.py` (the `python` name itself is
part of the check: python-is-python3). Creates a .docx, an .xlsx and a .pdf in a temp
dir, reads each one back with the same libraries and with poppler's pdftotext, and
exits non-zero on any mismatch. Prints one line per format.
"""
import os
import re
import shutil
import subprocess
import sys
import tempfile

assert os.path.realpath(sys.executable).startswith("/usr/bin/python3"), sys.executable  # Debian python3
assert sys.version_info[:2] >= (3, 11), sys.version

import docx  # python3-docx
import openpyxl  # python3-openpyxl
import pypdf  # python3-pypdf
from reportlab.lib.pagesizes import A4  # python3-reportlab
from reportlab.pdfgen import canvas

d = tempfile.mkdtemp(prefix="doc-tools-")
try:
    text = "WOOW pi-agent doc tools 2026"

    p = os.path.join(d, "t.docx")
    doc = docx.Document()
    doc.add_heading("Report", level=1)
    doc.add_paragraph(text)
    t = doc.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "a", "b"
    doc.save(p)
    back = docx.Document(p)
    assert any(text in para.text for para in back.paragraphs), "docx round-trip"
    assert back.tables[0].rows[0].cells[1].text == "b", "docx table"
    print("docx ok", os.path.getsize(p), "bytes, python-docx")

    p = os.path.join(d, "t.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["item", "qty"])
    ws.append(["x", 3])
    ws.append(["y", 4])
    ws["B4"] = "=SUM(B2:B3)"
    wb.save(p)
    ws2 = openpyxl.load_workbook(p).active
    assert ws2["A2"].value == "x" and ws2["B3"].value == 4 and ws2["B4"].value == "=SUM(B2:B3)", "xlsx round-trip"
    print("xlsx ok", os.path.getsize(p), "bytes, openpyxl", openpyxl.__version__)

    p = os.path.join(d, "t.pdf")
    c = canvas.Canvas(p, pagesize=A4)
    c.drawString(72, 760, text)
    c.showPage()
    c.drawString(72, 760, "page two")
    c.save()
    r = pypdf.PdfReader(p)
    assert len(r.pages) == 2 and text in r.pages[0].extract_text(), "pdf via pypdf"
    pdftotext = shutil.which("pdftotext")
    assert pdftotext, "poppler-utils pdftotext missing"
    out = subprocess.run([pdftotext, "-layout", p, "-"], check=True, capture_output=True, text=True).stdout
    assert text in out and "page two" in out, "pdf via pdftotext"
    info = subprocess.run(["pdfinfo", p], check=True, capture_output=True, text=True).stdout
    assert re.search(r"^Pages:\s+2$", info, re.M), info
    print("pdf ok", os.path.getsize(p), "bytes, reportlab + pypdf", pypdf.__version__, "+ poppler pdftotext/pdfinfo")

    # zh-TW PDF, the way a tenant's pi writes one: reportlab + an EMBEDDED TrueType CJK font
    # (the base's Noto CJK fonts are CFF, which reportlab refuses), read back by pypdf and
    # pdftotext, and pdffonts must show the font embedded.
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    zh = "渥屋科技 行政備忘錄：發票、報價單"
    pdfmetrics.registerFont(TTFont("WQY", "/usr/share/fonts/truetype/wqy/wqy-microhei.ttc", subfontIndex=0))
    p = os.path.join(d, "zh.pdf")
    c = canvas.Canvas(p, pagesize=A4)
    c.setFont("WQY", 16)
    c.drawString(72, 760, zh)
    c.save()
    assert zh in pypdf.PdfReader(p).pages[0].extract_text(), "zh pdf via pypdf"
    out = subprocess.run([pdftotext, p, "-"], check=True, capture_output=True, text=True).stdout
    assert zh in out, "zh pdf via pdftotext: %r" % out
    fonts = subprocess.run(["pdffonts", p], check=True, capture_output=True, text=True).stdout
    assert re.search(r"^\S*WenQuanYiMicroHei\S*\s+TrueType.*\byes\s+yes\s+yes\b", fonts, re.M), fonts
    print("zh pdf ok", os.path.getsize(p), "bytes, embedded WenQuanYi Micro Hei (pypdf + pdftotext + pdffonts)")
finally:
    shutil.rmtree(d, ignore_errors=True)
print("doc tools OK")
